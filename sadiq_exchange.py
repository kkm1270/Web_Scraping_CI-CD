#libraries
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook


#url
url="https://sadiqexchange.com/"
headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://www.google.com/",
    "Connection": "keep-alive"
}


#workbook
create_excel=Workbook()
create_sheet=create_excel.active
create_sheet.title="Currency Rates"
create_sheet.append(["Flag","Currency Rates"])

#load_website
rates=requests.get(url,headers=headers)
rates.raise_for_status()
parsing=BeautifulSoup(rates.text,"html.parser")
currencies=parsing('div',class_='ticker-item')

seen_currencies = set()

for currency in currencies:
    flag=currency.find('img')['src'] 
    text=currency.get_text(strip=True)
    if text not in seen_currencies:  # Check for duplicates
        seen_currencies.add(text)
        create_sheet.append([flag,text])
        print(flag,text)
    #create_sheet.append([flag,text])
    
    #print(flag,text)

create_excel.save("sadiq_exchange_scraping.xlsx")
print("Sadiq Exchange File Printed Successfully.")
